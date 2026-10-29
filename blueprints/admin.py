"""
Admin/back-office blueprint. Every route requires admin_required (least
privilege — cashiers get a generic 403 if they try any of these URLs
directly). This is also where Create/Read/Update/Delete for menu items and
staff accounts lives.
"""
from datetime import datetime, timedelta
from io import BytesIO
import os
import uuid

from flask import Blueprint, abort, current_app, flash, redirect, render_template, request, send_file, session, url_for
from openpyxl.chart import BarChart, Reference
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from werkzeug.security import generate_password_hash

from database.db import get_db, now_str
from database.product_descriptions import default_product_description
from security.audit import log_action
from security.decorators import admin_required
from security.validators import is_valid_email, is_valid_password, is_valid_pin, is_valid_price, is_valid_username

admin_bp = Blueprint("admin", __name__)

CATEGORIES = ("Coffee Drinks", "Freshie Drinks", "Meals", "Pastry")

AUDIT_CATEGORIES = {
    "Sales & Voids": {"SALE_COMPLETED", "VOID_APPROVED", "VOID_DENIED"},
    "Product Catalog": {
        "MENU_ITEM_CREATED", "MENU_ITEM_UPDATED", "MENU_ITEM_ARCHIVED",
        "MENU_ITEM_RESTORED", "MENU_ITEM_DELETED_PERMANENTLY",
        "MENU_IMAGE_UPDATED", "MENU_IMAGE_REMOVED",
    },
    "Staff Accounts": {
        "USER_CREATED", "STAFF_SIGNUP_REQUESTED", "STAFF_SIGNUP_APPROVED",
        "USER_DEACTIVATED", "PIN_RESET",
    },
    "Sign-in & Security": {
        "LOGIN_SUCCESS", "LOGIN_FAILED", "LOGOUT", "ACCOUNT_LOCKED",
    },
    "Password Recovery": {
        "PASSWORD_RESET_REQUESTED", "PASSWORD_RESET_FAILED", "PASSWORD_RESET_COMPLETED",
    },
}


def _excel_safe_values(values):
    safe_values = []
    for value in values:
        if isinstance(value, str) and value.lstrip(" \t\r\n")[:1] in ("=", "+", "-", "@"):
            value = "'" + value
        safe_values.append(value)
    return safe_values


def _remove_product_image(image_dir, filename):
    """Delete only a plain filename under the configured product image folder."""
    if not filename or os.path.basename(filename) != filename:
        return
    try:
        os.remove(os.path.join(image_dir, filename))
    except OSError:
        current_app.logger.warning("Could not remove product image %s", filename)


def _catalog_image_filename(upload):
    """Return a generated filename and validation error for an optional image."""
    if not upload or not upload.filename:
        return None, None
    header = upload.stream.read(12)
    upload.stream.seek(0)
    extension = os.path.splitext(upload.filename)[1].lower()
    signatures = {
        ".png": header.startswith(b"\x89PNG\r\n\x1a\n"),
        ".jpg": header.startswith(b"\xff\xd8\xff"),
        ".jpeg": header.startswith(b"\xff\xd8\xff"),
        ".webp": header.startswith(b"RIFF") and header[8:12] == b"WEBP",
    }
    if extension not in signatures or not signatures[extension]:
        return None, "Use a valid PNG, JPG, or WEBP image (maximum 5 MB)."
    return f"{uuid.uuid4().hex}{extension}", None


# ---------------------------------------------------------------- dashboard

@admin_bp.route("/dashboard")
@admin_required
def dashboard():
    db = get_db()
    period = request.args.get("period", "today")
    if period == "week":
        since_clause = "date(t.created_at) >= date('now','localtime','-6 days')"
    elif period == "month":
        since_clause = "date(t.created_at) >= date('now','localtime','start of month')"
    else:
        since_clause = "date(t.created_at) = date('now','localtime')"
        period = "today"

    summary = db.execute(
        f"SELECT COALESCE(SUM(total),0) AS revenue, COUNT(*) AS order_count "
        f"FROM transactions t WHERE status='completed' AND is_demo=0 AND {since_clause}"
    ).fetchone()

    # Profit uses the cost recorded on each line AT SALE TIME, so changing a
    # menu item's cost today never rewrites last week's profit figures.
    profit_row = db.execute(
        f"SELECT COALESCE(SUM((ti.unit_price - ti.unit_cost) * ti.quantity), 0) AS profit "
        f"FROM transaction_items ti JOIN transactions t ON t.id = ti.transaction_id "
        f"WHERE t.status = 'completed' AND t.is_demo=0 AND {since_clause}"
    ).fetchone()

    revenue = summary["revenue"] or 0
    profit = round(profit_row["profit"] or 0, 2)
    order_count = summary["order_count"] or 0
    margin = round(profit / revenue * 100, 1) if revenue else 0

    trending = db.execute(
        f"SELECT ti.item_name, SUM(ti.quantity) AS qty, SUM(ti.line_total) AS revenue "
        f"FROM transaction_items ti JOIN transactions t ON t.id = ti.transaction_id "
        f"WHERE t.status='completed' AND t.is_demo=0 AND {since_clause} "
        f"GROUP BY ti.item_name ORDER BY qty DESC LIMIT 5"
    ).fetchall()
    max_qty = max([t["qty"] for t in trending], default=1) or 1

    by_category = db.execute(
        f"SELECT mi.category, SUM(ti.line_total) AS revenue "
        f"FROM transaction_items ti "
        f"JOIN transactions t ON t.id = ti.transaction_id "
        f"JOIN menu_items mi ON mi.id = ti.menu_item_id "
        f"WHERE t.status='completed' AND t.is_demo=0 AND {since_clause} "
        f"GROUP BY mi.category ORDER BY revenue DESC"
    ).fetchall()
    max_cat = max([c["revenue"] for c in by_category], default=1) or 1

    payments = db.execute(
        f"SELECT payment_method, COUNT(*) AS n, COALESCE(SUM(total),0) AS amount "
        f"FROM transactions t WHERE status='completed' AND is_demo=0 AND {since_clause} "
        f"GROUP BY payment_method ORDER BY amount DESC"
    ).fetchall()

    last7 = db.execute(
        "SELECT date(created_at) AS d, COALESCE(SUM(total),0) AS revenue FROM transactions "
        "WHERE status='completed' AND is_demo=0 "
        "AND date(created_at) >= date('now','localtime','-6 days') "
        "GROUP BY date(created_at) ORDER BY d"
    ).fetchall()
    day_map = {r["d"]: r["revenue"] for r in last7}
    days = []
    for i in range(6, -1, -1):
        d = (datetime.now() - timedelta(days=i)).strftime("%Y-%m-%d")
        days.append({"date": d, "revenue": day_map.get(d, 0)})
    max_rev = max([d["revenue"] for d in days], default=1) or 1

    voided_count = db.execute(
        f"SELECT COUNT(*) AS c FROM transactions t "
        f"WHERE status='voided' AND is_demo=0 AND {since_clause}"
    ).fetchone()["c"]

    recent_txns = db.execute(
        "SELECT t.*, u.full_name AS cashier_name FROM transactions t "
        "JOIN users u ON u.id = t.cashier_id ORDER BY t.created_at DESC LIMIT 8"
    ).fetchall()

    pending_sync = db.execute(
        "SELECT COUNT(*) AS c FROM transactions WHERE sync_status='pending' AND is_demo=0"
    ).fetchone()["c"]

    return render_template(
        "admin_dashboard.html", period=period, revenue=revenue, profit=profit,
        order_count=order_count, margin=margin,
        trending=trending, max_qty=max_qty, by_category=by_category, max_cat=max_cat,
        payments=payments, days=days, max_rev=max_rev,
        voided_count=voided_count, recent_txns=recent_txns, pending_sync=pending_sync,
    )


# ---------------------------------------------------------------- menu (full CRUD)

@admin_bp.route("/menu")
@admin_required
def menu():
    db = get_db()
    items = db.execute(
        "SELECT * FROM menu_items WHERE is_archived = 0 ORDER BY "
        "CASE category WHEN 'Coffee Drinks' THEN 1 WHEN 'Freshie Drinks' THEN 2 "
        "WHEN 'Meals' THEN 3 WHEN 'Pastry' THEN 4 ELSE 5 END, name"
    ).fetchall()
    return render_template("admin_menu.html", items=items, categories=CATEGORIES, errors=None,
                           image_error=request.args.get("image_error"))


def _render_menu_with_errors(errors):
    db = get_db()
    items = db.execute(
        "SELECT * FROM menu_items WHERE is_archived = 0 ORDER BY "
        "CASE category WHEN 'Coffee Drinks' THEN 1 WHEN 'Freshie Drinks' THEN 2 "
        "WHEN 'Meals' THEN 3 WHEN 'Pastry' THEN 4 ELSE 5 END, name"
    ).fetchall()
    return render_template("admin_menu.html", items=items, categories=CATEGORIES, errors=errors)


@admin_bp.route("/menu/create", methods=["POST"])
@admin_required
def menu_create():
    name = (request.form.get("name") or "").strip()
    category = request.form.get("category")
    price = request.form.get("price")
    cost_price = request.form.get("cost_price")
    description = (request.form.get("description") or "").strip()
    image_upload = request.files.get("image")
    image_filename, image_error = _catalog_image_filename(image_upload)

    errors = []
    if not name or len(name) > 60:
        errors.append("Item name is required and must be 60 characters or fewer.")
    if len(description) > 180:
        errors.append("Product description must be 180 characters or fewer.")
    if category not in CATEGORIES:
        errors.append("Please choose a valid category.")
    if not is_valid_price(price):
        errors.append("Selling price must be a number greater than 0.")
    if not is_valid_price(cost_price):
        errors.append("Cost price must be a number greater than 0.")
    if image_error:
        errors.append(image_error)
    if request.content_length and request.content_length > current_app.config["MAX_CONTENT_LENGTH"]:
        errors.append("Image must be 5 MB or smaller.")
    if not errors and float(cost_price) > float(price):
        errors.append("Cost price cannot be higher than the selling price.")

    db = get_db()
    if not errors and db.execute("SELECT 1 FROM menu_items WHERE name = ?", (name,)).fetchone():
        errors.append("A menu item with that name already exists.")

    if errors:
        return _render_menu_with_errors(errors)

    price = round(float(price), 2)
    cost_price = round(float(cost_price), 2)
    description = description or default_product_description(name, category)
    image_dir = current_app.config["PRODUCT_IMAGE_DIR"]
    if image_filename:
        os.makedirs(image_dir, exist_ok=True)
        image_upload.save(os.path.join(image_dir, image_filename))
    try:
        db.execute(
            "INSERT INTO menu_items (name, category, price, cost_price, is_active, image_filename, description) "
            "VALUES (?, ?, ?, ?, 1, ?, ?)",
            (name, category, price, cost_price, image_filename, description),
        )
        db.commit()
    except Exception:
        if image_filename:
            _remove_product_image(image_dir, image_filename)
        raise
    log_action(db, session["user_id"], session.get("full_name"), "MENU_ITEM_CREATED",
               f"{name} ({category}) at \u20b1{price:.2f}, cost \u20b1{cost_price:.2f}")
    flash(f"Product {name} was added successfully.", "success")
    return redirect(url_for("admin.menu"))


@admin_bp.route("/menu/update", methods=["POST"])
@admin_required
def menu_update():
    item_id = request.form.get("item_id")
    price = request.form.get("price")
    cost_price = request.form.get("cost_price")
    description = (request.form.get("description") or "").strip()
    is_active = 1 if request.form.get("is_active") == "on" else 0

    db = get_db()
    item = db.execute("SELECT * FROM menu_items WHERE id = ?", (item_id,)).fetchone()
    if not item:
        abort(404)
    if len(description) > 180:
        return _render_menu_with_errors(["Product description must be 180 characters or fewer."])
    if not is_valid_price(price) or not is_valid_price(cost_price):
        return _render_menu_with_errors(["Price and cost must both be numbers greater than 0."])
    if float(cost_price) > float(price):
        return _render_menu_with_errors(["Cost price cannot be higher than the selling price."])

    price = round(float(price), 2)
    cost_price = round(float(cost_price), 2)
    description = description or default_product_description(item["name"], item["category"])
    db.execute(
        "UPDATE menu_items SET price = ?, cost_price = ?, is_active = ?, description = ? WHERE id = ?",
        (price, cost_price, is_active, description, item_id),
    )
    db.commit()
    log_action(db, session["user_id"], session.get("full_name"), "MENU_ITEM_UPDATED",
               f"{item['name']}: price \u20b1{item['price']:.2f} -> \u20b1{price:.2f}, "
               f"cost \u20b1{item['cost_price']:.2f} -> \u20b1{cost_price:.2f}, active={is_active}")
    return redirect(url_for("admin.menu"))


@admin_bp.route("/menu/<int:item_id>/remove", methods=["POST"])
@admin_required
def menu_remove(item_id):
    """Soft delete. The row stays so old receipts still resolve their item,
    it just leaves the terminal."""
    db = get_db()
    item = db.execute("SELECT * FROM menu_items WHERE id = ?", (item_id,)).fetchone()
    if not item:
        abort(404)
    db.execute(
        "UPDATE menu_items SET is_active = 0, is_archived = 1, archived_at = ? WHERE id = ?",
        (now_str(), item_id),
    )
    db.commit()
    log_action(db, session["user_id"], session.get("full_name"), "MENU_ITEM_ARCHIVED",
               f"Archived {item['name']} from the catalog")
    flash(f"{item['name']} was moved to the Archive.", "success")
    return redirect(url_for("admin.menu"))


@admin_bp.route("/menu/archive")
@admin_required
def menu_archive():
    db = get_db()
    items = db.execute(
        "SELECT m.*, (SELECT COUNT(*) FROM transaction_items ti WHERE ti.menu_item_id = m.id) AS sales_count "
        "FROM menu_items m WHERE m.is_archived = 1 ORDER BY m.archived_at DESC, m.name"
    ).fetchall()
    return render_template("admin_archive.html", items=items)


@admin_bp.route("/menu/<int:item_id>/restore", methods=["POST"])
@admin_required
def menu_restore(item_id):
    db = get_db()
    item = db.execute("SELECT * FROM menu_items WHERE id = ? AND is_archived = 1", (item_id,)).fetchone()
    if not item:
        abort(404)
    db.execute(
        "UPDATE menu_items SET is_archived = 0, archived_at = NULL, is_active = 1 WHERE id = ?",
        (item_id,),
    )
    db.commit()
    log_action(db, session["user_id"], session.get("full_name"), "MENU_ITEM_RESTORED",
               f"Restored {item['name']} to the catalog and active menu")
    flash(f"{item['name']} was restored to the active catalog.", "success")
    return redirect(url_for("admin.menu_archive"))


@admin_bp.route("/menu/<int:item_id>/delete", methods=["POST"])
@admin_required
def menu_delete_permanently(item_id):
    db = get_db()
    item = db.execute(
        "SELECT * FROM menu_items WHERE id = ? AND is_archived = 1", (item_id,)
    ).fetchone()
    if not item:
        abort(404)
    sales_count = db.execute(
        "SELECT COUNT(*) AS c FROM transaction_items WHERE menu_item_id = ?", (item_id,)
    ).fetchone()["c"]
    if sales_count:
        flash(
            f"{item['name']} has sales history and cannot be permanently deleted. It remains archived.",
            "warning",
        )
        return redirect(url_for("admin.menu_archive"))

    db.execute("DELETE FROM menu_items WHERE id = ?", (item_id,))
    db.commit()
    log_action(db, session["user_id"], session.get("full_name"), "MENU_ITEM_DELETED_PERMANENTLY",
               f"Permanently deleted unused archived product {item['name']}")
    _remove_product_image(current_app.config["PRODUCT_IMAGE_DIR"], item["image_filename"])
    flash(f"{item['name']} was permanently deleted.", "success")
    return redirect(url_for("admin.menu_archive"))


@admin_bp.route("/menu/<int:item_id>/image", methods=["POST"])
@admin_required
def menu_image_upload(item_id):
    """Add or replace a catalog image after checking its actual file signature."""
    db = get_db()
    item = db.execute("SELECT * FROM menu_items WHERE id = ?", (item_id,)).fetchone()
    if not item:
        abort(404)

    upload = request.files.get("image")
    if not upload or not upload.filename:
        return redirect(url_for("admin.menu", image_error="Choose an image file first."))
    if request.content_length and request.content_length > current_app.config["MAX_CONTENT_LENGTH"]:
        return redirect(url_for("admin.menu", image_error="Image must be 5 MB or smaller."))

    filename, validation_error = _catalog_image_filename(upload)
    if validation_error:
        return redirect(url_for("admin.menu", image_error=validation_error))
    image_dir = current_app.config["PRODUCT_IMAGE_DIR"]
    os.makedirs(image_dir, exist_ok=True)
    upload.save(os.path.join(image_dir, filename))
    try:
        db.execute("UPDATE menu_items SET image_filename = ? WHERE id = ?", (filename, item_id))
        db.commit()
        log_action(db, session["user_id"], session.get("full_name"), "MENU_IMAGE_UPDATED",
                   f"Updated product image for {item['name']}")
    except Exception:
        try:
            os.remove(os.path.join(image_dir, filename))
        except OSError:
            pass
        raise

    old_filename = item["image_filename"]
    if old_filename and old_filename != filename:
        _remove_product_image(image_dir, old_filename)
    flash(f"Image for {item['name']} was uploaded successfully.", "success")
    return redirect(url_for("admin.menu"))


@admin_bp.route("/menu/<int:item_id>/image/remove", methods=["POST"])
@admin_required
def menu_image_remove(item_id):
    db = get_db()
    item = db.execute("SELECT * FROM menu_items WHERE id = ?", (item_id,)).fetchone()
    if not item:
        abort(404)
    filename = item["image_filename"]
    if filename:
        db.execute("UPDATE menu_items SET image_filename = NULL WHERE id = ?", (item_id,))
        db.commit()
        log_action(db, session["user_id"], session.get("full_name"), "MENU_IMAGE_REMOVED",
                   f"Removed product image for {item['name']}")
        _remove_product_image(current_app.config["PRODUCT_IMAGE_DIR"], filename)
        flash(f"Image for {item['name']} was removed.", "success")
    return redirect(url_for("admin.menu"))


# ---------------------------------------------------------------- users (CRUD)

@admin_bp.route("/users")
@admin_required
def users():
    db = get_db()
    all_users = db.execute(
        "SELECT * FROM users ORDER BY "
        "CASE approval_status WHEN 'pending' THEN 0 ELSE 1 END, role, full_name"
    ).fetchall()
    return render_template("admin_users.html", users=all_users, errors=None)


@admin_bp.route("/users/create", methods=["POST"])
@admin_required
def users_create():
    full_name = (request.form.get("full_name") or "").strip()
    role = request.form.get("role")
    pin = request.form.get("pin")

    errors = []
    if not full_name:
        errors.append("Full name is required.")
    if role not in ("admin", "cashier"):
        errors.append("Invalid role.")
    if not is_valid_pin(pin):
        errors.append("PIN must be exactly 4 digits and not an easily-guessed pattern (e.g. 1234, 0000).")

    # Admins use an email and password; staff use a username and PIN, with
    # an optional email as an alternate sign-in identifier.
    if role == "admin":
        username = email = (request.form.get("email") or "").strip().lower()
        password = request.form.get("password")
        if not is_valid_email(email):
            errors.append("Enter a valid email address.")
        if not is_valid_password(password):
            errors.append("Admin accounts need a password: 8+ characters with a letter and a number.")
    else:
        username = (request.form.get("username") or "").strip()
        email = (request.form.get("email") or "").strip().lower() or None
        password = None
        if not is_valid_username(username):
            errors.append("Staff ID must be 3-20 characters: letters, numbers, underscore.")
        if email and not is_valid_email(email):
            errors.append("Enter a valid email address, or leave it blank.")

    db = get_db()
    if not errors:
        taken = db.execute(
            "SELECT 1 FROM users WHERE lower(username) IN (?, ?) "
            "OR lower(email) IN (?, ?) LIMIT 1",
            (username.lower(), (email or "").lower(), username.lower(), (email or "").lower()),
        ).fetchone()
        if taken:
            errors.append("An account with that email already exists." if role == "admin"
                          else "That Staff ID is already taken.")

    if errors:
        all_users = db.execute("SELECT * FROM users ORDER BY role, full_name").fetchall()
        return render_template("admin_users.html", users=all_users, errors=errors)

    password_hash = generate_password_hash(password) if role == "admin" else None
    db.execute(
        "INSERT INTO users "
        "(username, email, password_hash, pin_hash, full_name, role, approval_status, is_active, created_at) "
        "VALUES (?, ?, ?, ?, ?, ?, 'approved', 1, ?)",
        (username, email, password_hash, generate_password_hash(pin), full_name, role, now_str()),
    )
    db.commit()
    log_action(db, session["user_id"], session.get("full_name"), "USER_CREATED",
               f"Created {role} account for {full_name} ({email if role == 'admin' else username})")
    return redirect(url_for("admin.users"))


@admin_bp.route("/users/<int:user_id>/approve", methods=["POST"])
@admin_required
def users_approve(user_id):
    db = get_db()
    user = db.execute(
        "SELECT * FROM users WHERE id = ? AND role = 'cashier' AND approval_status = 'pending'",
        (user_id,),
    ).fetchone()
    if not user:
        abort(404)

    db.execute(
        "UPDATE users SET approval_status = 'approved', is_active = 1, "
        "failed_attempts = 0, locked_until = NULL WHERE id = ?",
        (user_id,),
    )
    db.commit()
    log_action(db, session["user_id"], session.get("full_name"), "STAFF_SIGNUP_APPROVED",
               f"Approved staff account request for {user['username']}")
    return redirect(url_for("admin.users"))


@admin_bp.route("/users/<int:user_id>/deactivate", methods=["POST"])
@admin_required
def users_deactivate(user_id):
    db = get_db()
    user = db.execute("SELECT * FROM users WHERE id = ?", (user_id,)).fetchone()
    if not user:
        abort(404)
    if user_id == session["user_id"]:
        return redirect(url_for("admin.users"))  # can't deactivate your own account
    db.execute("UPDATE users SET is_active = 0 WHERE id = ?", (user_id,))
    db.commit()
    log_action(db, session["user_id"], session.get("full_name"), "USER_DEACTIVATED",
               f"Deactivated {user['full_name']} ({user['username']})")
    return redirect(url_for("admin.users"))


@admin_bp.route("/users/<int:user_id>/reset-pin", methods=["POST"])
@admin_required
def users_reset_pin(user_id):
    new_pin = request.form.get("new_pin")
    db = get_db()
    user = db.execute("SELECT * FROM users WHERE id = ?", (user_id,)).fetchone()
    if not user:
        abort(404)
    if not is_valid_pin(new_pin):
        return redirect(url_for("admin.users"))
    db.execute(
        "UPDATE users SET pin_hash = ?, failed_attempts = 0, locked_until = NULL WHERE id = ?",
        (generate_password_hash(new_pin), user_id),
    )
    db.commit()
    log_action(db, session["user_id"], session.get("full_name"), "PIN_RESET", f"Reset PIN for {user['full_name']}")
    return redirect(url_for("admin.users"))


# ---------------------------------------------------------------- audit log & transactions (Read)

@admin_bp.route("/audit-log")
@admin_required
def audit_log():
    db = get_db()
    records = db.execute("SELECT * FROM audit_log ORDER BY created_at DESC LIMIT 500").fetchall()
    all_logs = []
    category_counts = {name: 0 for name in AUDIT_CATEGORIES}
    category_counts["Other"] = 0
    for record in records:
        log = dict(record)
        log["category"] = next(
            (name for name, actions in AUDIT_CATEGORIES.items() if log["action"] in actions),
            "Other",
        )
        category_counts[log["category"]] += 1
        all_logs.append(log)

    selected_category = request.args.get("category", "All")
    if selected_category not in ("All", *category_counts.keys()):
        selected_category = "All"
    logs = [log for log in all_logs
            if selected_category == "All" or log["category"] == selected_category]
    return render_template(
        "admin_audit_log.html", logs=logs, audit_categories=category_counts,
        selected_category=selected_category, total_count=len(all_logs),
    )


@admin_bp.route("/transactions")
@admin_required
def transactions():
    db = get_db()
    txns = db.execute(
        "SELECT t.*, u.full_name AS cashier_name, v.full_name AS voided_by_name "
        "FROM transactions t JOIN users u ON u.id = t.cashier_id "
        "LEFT JOIN users v ON v.id = t.voided_by ORDER BY t.created_at DESC LIMIT 100"
    ).fetchall()
    return render_template("admin_transactions.html", txns=txns)


@admin_bp.route("/transactions/export.xlsx")
@admin_required
def transactions_export():
    db = get_db()
    today = datetime.now().date()
    period = request.args.get("period", "all")
    period_names = {"all": "All time", "week": "This week", "month": "This month", "year": "This year"}
    if period not in period_names:
        period = "all"
    start_date = end_date = None
    if period == "week":
        start_date = today - timedelta(days=today.weekday())
        end_date = today + timedelta(days=1)
    elif period == "month":
        start_date = today.replace(day=1)
        end_date = today + timedelta(days=1)
    elif period == "year":
        start_date = today.replace(month=1, day=1)
        end_date = today + timedelta(days=1)

    date_clause = ""
    date_params = ()
    if start_date is not None:
        date_clause = " AND t.created_at >= ? AND t.created_at < ?"
        date_params = (
            start_date.strftime("%Y-%m-%d 00:00:00"),
            end_date.strftime("%Y-%m-%d 00:00:00"),
        )

    summary = db.execute(
        "SELECT COUNT(*) AS transaction_count, COALESCE(SUM(total), 0) AS revenue "
        "FROM transactions t WHERE status = 'completed' AND is_demo = 0" + date_clause,
        date_params,
    ).fetchone()
    top_products = db.execute(
        "SELECT ti.item_name, SUM(ti.quantity) AS quantity "
        "FROM transaction_items ti JOIN transactions t ON t.id = ti.transaction_id "
        "WHERE t.status = 'completed' AND t.is_demo = 0" + date_clause +
        " GROUP BY ti.item_name ORDER BY quantity DESC, ti.item_name LIMIT 5",
        date_params,
    ).fetchall()
    transactions = db.execute(
        "SELECT t.*, cashier.full_name AS cashier_name, voider.full_name AS voided_by_name "
        "FROM transactions t "
        "JOIN users cashier ON cashier.id = t.cashier_id "
        "LEFT JOIN users voider ON voider.id = t.voided_by WHERE 1=1" + date_clause +
        " ORDER BY t.created_at DESC, t.id DESC",
        date_params,
    ).fetchall()
    items = db.execute(
        "SELECT ti.transaction_id, t.created_at, cashier.full_name AS cashier_name, "
        "t.service_type, t.payment_method, t.status, t.is_demo, ti.item_name, "
        "ti.unit_price, ti.unit_cost, ti.quantity, ti.line_total "
        "FROM transaction_items ti "
        "JOIN transactions t ON t.id = ti.transaction_id "
        "JOIN users cashier ON cashier.id = t.cashier_id "
        "WHERE 1=1" + date_clause + " ORDER BY ti.transaction_id DESC, ti.id",
        date_params,
    ).fetchall()
    menu_products = db.execute(
        "SELECT id, name, category, price, cost_price, is_active "
        "FROM menu_items ORDER BY "
        "CASE category WHEN 'Coffee Drinks' THEN 1 WHEN 'Freshie Drinks' THEN 2 "
        "WHEN 'Meals' THEN 3 WHEN 'Pastry' THEN 4 ELSE 5 END, name"
    ).fetchall()
    items_by_transaction = {}
    for item in items:
        items_by_transaction.setdefault(item["transaction_id"], []).append(
            f"{item['item_name']} x {item['quantity']}"
        )

    workbook = Workbook()
    dashboard = workbook.active
    dashboard.title = "Dashboard"
    dashboard.sheet_view.showGridLines = False
    dashboard.merge_cells("A1:N1")
    dashboard["A1"] = "Cafe Sales Dashboard"
    dashboard["A1"].font = Font(bold=True, size=20, color="FFFFFF")
    dashboard["A1"].fill = PatternFill("solid", fgColor="4A3226")
    dashboard["A1"].alignment = Alignment(vertical="center")
    dashboard.row_dimensions[1].height = 36
    dashboard.merge_cells("A2:N2")
    dashboard["A2"] = (
        f"{period_names[period]}; summary excludes demo and voided sales, transaction chart shows all statuses "
        f"| Generated {today:%B %d, %Y}"
    )
    dashboard["A2"].font = Font(italic=True, color="6B5B53")

    for label_range, value_range, label, value, number_format in (
        ("A4:F4", "A5:F6", "COMPLETED TRANSACTIONS", summary["transaction_count"], "#,##0"),
        ("H4:M4", "H5:M6", "TOTAL REVENUE", summary["revenue"], '"₱"#,##0.00'),
    ):
        dashboard.merge_cells(label_range)
        dashboard.merge_cells(value_range)
        label_cell = dashboard[label_range.split(":")[0]]
        value_cell = dashboard[value_range.split(":")[0]]
        label_cell.value = label
        label_cell.font = Font(bold=True, color="FFFFFF", size=11)
        label_cell.fill = PatternFill("solid", fgColor="80644F")
        label_cell.alignment = Alignment(horizontal="center", vertical="center")
        value_cell.value = value
        value_cell.font = Font(bold=True, color="4A3226", size=24)
        value_cell.fill = PatternFill("solid", fgColor="F3E9DF")
        value_cell.alignment = Alignment(horizontal="center", vertical="center")
        value_cell.number_format = number_format
    dashboard.row_dimensions[4].height = 25
    dashboard.row_dimensions[5].height = 30
    dashboard.row_dimensions[6].height = 20
    for column in range(1, 15):
        dashboard.column_dimensions[get_column_letter(column)].width = 13

    chart_data = workbook.create_sheet("Chart Data")
    chart_data.append(("Transaction", "Transaction Total (PHP)", "Product", "Quantity Sold"))
    for index, row in enumerate(transactions):
        status_label = row["status"]
        if row["is_demo"]:
            status_label = f"{status_label}, demo"
        label = f"#{row['id']} | {row['created_at'][:10]} | {status_label}"
        chart_data.append((
            label,
            row["total"],
            top_products[index]["item_name"] if index < len(top_products) else None,
            top_products[index]["quantity"] if index < len(top_products) else None,
        ))
    if not transactions:
        chart_data.append((None, None, None, None))

    revenue_chart = BarChart()
    revenue_chart.type = "col"
    revenue_chart.style = 10
    revenue_chart.title = "All Transactions by Sale"
    revenue_chart.y_axis.title = "Transaction Total (PHP)"
    revenue_chart.x_axis.title = "Transaction | Date | Status"
    revenue_chart.height = 8
    revenue_chart.width = 15
    revenue_chart.legend = None
    last_transaction_row = len(transactions) + 1 if transactions else 2
    revenue_chart.add_data(
        Reference(chart_data, min_col=2, min_row=1, max_row=last_transaction_row),
        titles_from_data=True,
    )
    revenue_chart.set_categories(
        Reference(chart_data, min_col=1, min_row=2, max_row=last_transaction_row)
    )
    revenue_chart.y_axis.numFmt = '"₱"#,##0'
    dashboard.add_chart(revenue_chart, "A8")

    products_chart = BarChart()
    products_chart.type = "bar"
    products_chart.style = 11
    products_chart.title = "Top 5 Most Ordered Products"
    products_chart.x_axis.title = "Quantity Sold"
    products_chart.y_axis.title = "Product"
    products_chart.height = 8
    products_chart.width = 15
    products_chart.legend = None
    products_chart.add_data(Reference(chart_data, min_col=4, min_row=1, max_row=max(2, len(top_products) + 1)),
                            titles_from_data=True)
    products_chart.set_categories(Reference(chart_data, min_col=3, min_row=2, max_row=max(2, len(top_products) + 1)))
    dashboard.add_chart(products_chart, "I8")
    chart_data.sheet_state = "hidden"

    transaction_sheet = workbook.create_sheet("Transactions")
    transaction_sheet.title = "Transactions"
    transaction_headers = (
        "Transaction ID", "Date", "Cashier", "Order Type", "Payment Method",
        "Subtotal", "VAT", "Total", "Amount Tendered", "Change",
        "Status", "Demo Sale", "Voided By", "Void Reason", "Backup Status",
        "Products Sold",
    )
    transaction_sheet.append(transaction_headers)
    for row in transactions:
        transaction_sheet.append(_excel_safe_values((
            row["id"], row["created_at"], row["cashier_name"], row["service_type"],
            row["payment_method"], row["subtotal"], row["vat_amount"], row["total"],
            row["amount_tendered"], row["change_due"], row["status"],
            "Yes" if row["is_demo"] else "No", row["voided_by_name"],
            row["void_reason"], row["sync_status"],
            "; ".join(items_by_transaction.get(row["id"], [])),
        )))

    item_sheet = workbook.create_sheet("Transaction Items")
    item_sheet.append((
        "Transaction ID", "Date", "Cashier", "Order Type", "Payment Method",
        "Status", "Demo Sale", "Item", "Unit Price", "Unit Cost", "Quantity", "Line Total",
    ))
    for row in items:
        item_sheet.append(_excel_safe_values((
            row["transaction_id"], row["created_at"], row["cashier_name"],
            row["service_type"], row["payment_method"], row["status"],
            "Yes" if row["is_demo"] else "No", row["item_name"], row["unit_price"],
            row["unit_cost"], row["quantity"], row["line_total"],
        )))

    product_sheet = workbook.create_sheet("Menu Products")
    product_sheet.append((
        "Product ID", "Product", "Category", "Selling Price", "Unit Cost", "Menu Status",
    ))
    for row in menu_products:
        product_sheet.append(_excel_safe_values((
            row["id"], row["name"], row["category"], row["price"], row["cost_price"],
            "Active" if row["is_active"] else "Hidden",
        )))

    for sheet in (transaction_sheet, item_sheet, product_sheet):
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.row_dimensions[1].height = 24
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="4A3226")
        currency_columns = (
            (6, 7, 8, 9, 10) if sheet is transaction_sheet else
            (9, 10, 12) if sheet is item_sheet else
            (4, 5)
        )
        for row_number in range(2, sheet.max_row + 1):
            for column_number in currency_columns:
                sheet.cell(row=row_number, column=column_number).number_format = '"₱"#,##0.00'
        for column in sheet.columns:
            values = [len(str(cell.value or "")) for cell in column]
            sheet.column_dimensions[get_column_letter(column[0].column)].width = min(
                max(max(values, default=10) + 2, 12), 34
            )
    transaction_sheet.column_dimensions["P"].width = 48
    for row_number in range(2, transaction_sheet.max_row + 1):
        transaction_sheet.cell(row=row_number, column=16).alignment = Alignment(
            vertical="top", wrap_text=True
        )

    # Keep the export focused on transaction data; these sheets only support
    # the optional overview and catalog tabs requested to be removed.
    for sheet_name in ("Dashboard", "Chart Data", "Menu Products"):
        if sheet_name in workbook.sheetnames:
            workbook.remove(workbook[sheet_name])
    workbook.active = workbook.sheetnames.index("Transactions")

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return send_file(
        output,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=f"cafe-sales-{period}-{datetime.now():%Y-%m-%d}.xlsx",
    )
