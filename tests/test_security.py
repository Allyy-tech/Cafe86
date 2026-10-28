"""
Automated security tests.

Each test maps to a claim made in the project documentation, so you can
demonstrate that the security measures actually work rather than just
describing them. Run with:

    python tests/test_security.py
"""
import os
import re
import sys
import tempfile
import unittest
from io import BytesIO
from unittest.mock import patch

from flask.testing import FlaskClient

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from app import create_app  # noqa: E402
from database.db import get_db  # noqa: E402


class CaptchaTestClient(FlaskClient):
    """Solve the visible addition CAPTCHA for auth requests in integration tests."""

    def post(self, *args, **kwargs):
        path = args[0].split("?", 1)[0] if args else ""
        if path in ("/admin-login", "/pos-login", "/", "/signup"):
            page_path = "/signup" if path == "/signup" else "/admin-login"
            page = super().get(page_path)
            match = re.search(rb"Security check: (\d+) \+ (\d+) = \?", page.data)
            if match:
                answer = str(int(match.group(1)) + int(match.group(2)))
                data = dict(kwargs.get("data") or {})
                data.setdefault("captcha_answer", answer)
                kwargs["data"] = data
        return super().post(*args, **kwargs)


class BaseCase(unittest.TestCase):
    # Every existing test assumes an admin already exists (email
    # 'admin@example.com', password 'Admin#2026', pin '7391'). seed_data.py
    # no longer creates one — the first admin comes from the /admin-setup
    # bootstrap flow — so this fixture inserts that same account directly
    # via SQL instead, keeping every other test unchanged. Both username and
    # email are set to the same value here, matching what a real admin
    # account looks like (see auth.py's admin_setup / admin.py's
    # users_create, which both do the same thing).
    # TestBootstrapAdminSetup sets seed_admin False to get a database with
    # genuinely zero admins.
    seed_admin = True

    def setUp(self):
        self.tmpdir = tempfile.mkdtemp()
        os.environ["SECRET_KEY"] = "test-secret-key-not-for-production"
        self.app = create_app()
        self.app.config["TESTING"] = True
        self.app.config["DATABASE_PATH"] = os.path.join(self.tmpdir, "test.db")
        self.app.test_client_class = CaptchaTestClient

        with self.app.app_context():
            from database.seed_data import ensure_seed_data, init_schema
            init_schema(self.app)
            ensure_seed_data()
            if self.seed_admin:
                self._seed_test_admin()

        self.client = self.app.test_client()

    def _seed_test_admin(self):
        from werkzeug.security import generate_password_hash
        db = get_db()
        db.execute(
            "INSERT INTO users (username, email, password_hash, pin_hash, full_name, role, is_active, created_at) "
            "VALUES ('admin@example.com', 'admin@example.com', ?, ?, 'Admin Account', 'admin', 1, datetime('now'))",
            (generate_password_hash("Admin#2026"), generate_password_hash("7391")),
        )
        db.commit()

    def csrf(self):
        return self.client.get("/api/csrf-token").get_json()["token"]

    def login_admin(self):
        return self.client.post("/", data={
            "identifier": "admin@example.com", "credential": "Admin#2026", "csrf_token": self.csrf(),
        }, follow_redirects=True)

    def login_cashier(self):
        return self.client.post("/", data={
            "identifier": "cashier1", "credential": "1428", "csrf_token": self.csrf(),
        }, follow_redirects=True)

    def sell(self, menu_item_id=1, quantity=1, method="cash", tendered=1000):
        return self.client.post("/pos/checkout", json={
            "cart": [{"menu_item_id": menu_item_id, "quantity": quantity}],
            "payment_method": method, "service_type": "dine_in",
            "amount_tendered": tendered,
        }, headers={"X-CSRF-Token": self.csrf()})


class TestAuthentication(BaseCase):
    """IA Principle: Authentication."""

    def test_landing_has_unified_sign_in_captcha_and_staff_signup(self):
        res = self.client.get("/")
        self.assertEqual(res.status_code, 200)
        self.assertIn(b"Username or email", res.data)
        self.assertIn(b"Admin password or staff PIN", res.data)
        self.assertIn(b"Security check:", res.data)
        self.assertIn(b"Request an account", res.data)
        self.assertIn(b"Forgot password?", res.data)
        self.assertIn(b'action="/"', res.data)

    def test_passwords_are_hashed_not_plaintext(self):
        with self.app.app_context():
            row = get_db().execute("SELECT password_hash, pin_hash FROM users WHERE email='admin@example.com'").fetchone()
        self.assertNotIn("Admin#2026", row["password_hash"])
        self.assertNotIn("7391", row["pin_hash"])
        self.assertTrue(row["password_hash"].startswith(("pbkdf2:", "scrypt:")))

    def test_valid_admin_login_succeeds(self):
        res = self.login_admin()
        self.assertEqual(res.status_code, 200)
        self.assertIn(b"Owner Dashboard", res.data)

    def test_admin_forgot_password_remains_available(self):
        res = self.client.get("/admin-login")
        self.assertIn(b"Forgot password?", res.data)
        self.assertIn(b"href=\"/forgot-password\"", res.data)

    def test_wrong_password_is_rejected(self):
        res = self.client.post("/admin-login", data={
            "email": "admin@example.com", "password": "wrongpassword", "csrf_token": self.csrf()})
        self.assertIn(b"Invalid username/email or password/PIN", res.data)

    def test_error_message_does_not_reveal_whether_user_exists(self):
        real = self.client.post("/admin-login", data={
            "email": "admin@example.com", "password": "wrong1234", "csrf_token": self.csrf()}).data
        fake = self.client.post("/admin-login", data={
            "email": "ghost@example.com", "password": "wrong1234", "csrf_token": self.csrf()}).data
        self.assertIn(b"Invalid username/email or password/PIN", real)
        self.assertIn(b"Invalid username/email or password/PIN", fake)

    def test_account_locks_after_repeated_failures(self):
        for _ in range(5):
            self.client.post("/admin-login", data={
                "email": "admin@example.com", "password": "badpass123", "csrf_token": self.csrf()})
        res = self.client.post("/admin-login", data={
            "email": "admin@example.com", "password": "Admin#2026", "csrf_token": self.csrf()})
        self.assertIn(b"Unable to sign in", res.data)
        with self.app.app_context():
            locked_until = get_db().execute(
                "SELECT locked_until FROM users WHERE email='admin@example.com'"
            ).fetchone()["locked_until"]
        self.assertIsNotNone(locked_until)

    def test_cashier_pin_login_works(self):
        res = self.login_cashier()
        self.assertEqual(res.status_code, 200)
        self.assertIn(b"<h1>Menu</h1>", res.data)

    def test_staff_id_and_pin_are_both_required(self):
        wrong_id = self.client.post("/pos-login", data={
            "staff_id": "unknown", "pin": "1428", "csrf_token": self.csrf(),
        })
        wrong_pin = self.client.post("/pos-login", data={
            "staff_id": "cashier1", "pin": "9999", "csrf_token": self.csrf(),
        })
        self.assertIn(b"Invalid username/email or password/PIN", wrong_id.data)
        self.assertIn(b"Invalid username/email or password/PIN", wrong_pin.data)
        self.assertNotIn(b"New Order", wrong_id.data)
        self.assertNotIn(b"New Order", wrong_pin.data)

    def test_staff_sign_in_page_does_not_list_employee_names(self):
        res = self.client.get("/pos-login")
        self.assertEqual(res.status_code, 200)
        self.assertIn(b"Username or email", res.data)
        self.assertIn(b"Staff use their private PIN", res.data)
        self.assertNotIn(b"Cashier One", res.data)

    def test_incorrect_captcha_prevents_login(self):
        res = self.client.post("/admin-login", data={
            "identifier": "admin@example.com", "credential": "Admin#2026",
            "captcha_answer": "0", "csrf_token": self.csrf(),
        })
        self.assertIn(b"addition answer is incorrect", res.data)
        self.assertNotIn(b"Owner Dashboard", res.data)

    def test_staff_signup_waits_for_admin_approval(self):
        signup = self.client.post("/signup", data={
            "full_name": "New Barista", "username": "newbarista",
            "email": "barista@example.com", "pin": "5839", "role": "admin",
            "csrf_token": self.csrf(),
        })
        self.assertIn(b"Request submitted", signup.data)
        with self.app.app_context():
            user = get_db().execute(
                "SELECT id, role, approval_status, is_active FROM users WHERE username='newbarista'"
            ).fetchone()
        self.assertEqual((user["role"], user["approval_status"], user["is_active"]),
                         ("cashier", "pending", 0))

        denied = self.client.post("/pos-login", data={
            "identifier": "newbarista", "credential": "5839", "csrf_token": self.csrf(),
        })
        self.assertIn(b"Invalid username/email or password/PIN", denied.data)

        self.login_admin()
        approved = self.client.post(
            f"/admin/users/{user['id']}/approve", data={"csrf_token": self.csrf()}
        )
        self.assertEqual(approved.status_code, 302)
        self.client.get("/logout")
        allowed = self.client.post("/pos-login", data={
            "identifier": "barista@example.com", "credential": "5839",
            "csrf_token": self.csrf(),
        }, follow_redirects=True)
        self.assertIn(b"<h1>Menu</h1>", allowed.data)

    def test_admin_dashboard_uses_requested_labels_and_hides_removed_metrics(self):
        self.login_admin()
        res = self.client.get("/admin/dashboard")
        self.assertIn(b"Most Ordered", res.data)
        self.assertNotIn(b"Average Order", res.data)
        self.assertNotIn(b"Sales by Cashier", res.data)

    def test_today_sales_has_five_labeled_demo_transactions(self):
        self.login_admin()
        from datetime import datetime
        just_after_midnight = datetime.now().replace(hour=0, minute=5, second=0, microsecond=0)
        with patch("database.seed_data.datetime") as mock_datetime:
            mock_datetime.now.return_value = just_after_midnight
            res = self.client.get("/pos/history")
        self.assertEqual(res.status_code, 200)
        self.assertEqual(res.data.count(b"DEMO SALE"), 5)
        self.assertIn(b"0.00", res.data)
        dashboard = self.client.get("/admin/dashboard")
        self.assertIn(b"0 completed order(s)", dashboard.data)
        with self.app.app_context():
            db = get_db()
            row = db.execute(
                "SELECT COUNT(*) AS count, COALESCE(SUM(total), 0) AS total "
                "FROM transactions WHERE is_demo = 1"
            ).fetchone()
            distinct_dates = db.execute(
                "SELECT COUNT(DISTINCT date(created_at)) FROM transactions WHERE is_demo = 1"
            ).fetchone()[0]
        self.assertEqual(row["count"], 5)
        self.assertEqual(distinct_dates, 1)
        self.assertGreater(row["total"], 0)

    def test_pos_menu_shows_grouped_products_and_payment_options(self):
        self.login_cashier()
        res = self.client.get("/pos/")
        self.assertEqual(res.status_code, 200)
        for text in (b"Coffee Drinks", b"Freshie Drinks", b"merchant-qr",
                     b"Dine In", b"Take Out", b"GCash"):
            self.assertIn(text, res.data)

    def test_merchant_qr_requires_login_and_is_not_public_in_demo_mode(self):
        self.assertEqual(self.client.get("/pos/merchant-qr").status_code, 302)
        self.login_cashier()
        qr = self.client.get("/pos/merchant-qr")
        self.assertEqual(qr.status_code, 200)
        self.assertEqual(qr.mimetype, "image/jpeg")

        self.app.config["PUBLIC_DEMO"] = True
        terminal = self.client.get("/pos/")
        self.assertIn(b"Demo unavailable", terminal.data)
        self.assertNotIn(b"/pos/merchant-qr", terminal.data)
        self.assertEqual(self.client.get("/pos/merchant-qr").status_code, 404)

    def test_public_demo_rejects_ewallet_checkout(self):
        self.login_cashier()
        self.app.config["PUBLIC_DEMO"] = True
        response = self.sell(method="ewallet")
        self.assertEqual(response.status_code, 400)
        self.assertIn(b"E-wallet payments are disabled", response.data)

    def test_admin_can_export_transactions_and_line_items_as_xlsx(self):
        self.login_cashier()
        self.sell()
        self.sell(menu_item_id=2)
        with self.app.app_context():
            get_db().execute("UPDATE transactions SET status='voided' WHERE id=1")
            get_db().commit()
        self.client.get("/logout")
        self.login_admin()

        response = self.client.get("/admin/transactions/export.xlsx")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response.mimetype,
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        from openpyxl import load_workbook
        workbook = load_workbook(BytesIO(response.data))
        self.assertEqual(
            workbook.sheetnames,
            ["Dashboard", "Chart Data", "Transactions", "Transaction Items", "Menu Products"],
        )
        dashboard = workbook["Dashboard"]
        self.assertEqual(dashboard["A5"].value, 1)
        self.assertEqual(dashboard["H5"].value, workbook["Transactions"]["H2"].value)
        self.assertEqual(len(dashboard._charts), 2)
        self.assertEqual(
            [chart.title.tx.rich.p[0].r[0].t for chart in dashboard._charts],
            ["All Transactions by Sale", "Top 5 Most Ordered Products"],
        )
        chart_data = workbook["Chart Data"]
        self.assertEqual(chart_data.max_row, workbook["Transactions"].max_row)
        self.assertIn("#2 |", chart_data["A2"].value)
        self.assertIn("#1 |", chart_data["A3"].value)
        self.assertIn("voided", chart_data["A3"].value)
        self.assertEqual(dashboard._charts[0].series[0].val.numRef.f, "'Chart Data'!$B$2:$B$3")
        self.assertEqual(workbook["Chart Data"].sheet_state, "hidden")
        self.assertEqual(workbook["Transactions"]["A2"].value, 2)
        self.assertEqual(workbook["Transactions"]["P2"].value, "Americano x 1")
        self.assertEqual(workbook["Transactions"]["P3"].value, "Espresso x 1")
        self.assertEqual(workbook["Transaction Items"]["A2"].value, 2)
        self.assertEqual(workbook["Menu Products"]["B2"].value, "Americano")
        self.assertEqual(workbook["Menu Products"].max_row, 17)

    def test_deployment_mode_seeds_menu_without_demo_staff_accounts(self):
        from app import Config, create_app
        database_path = os.path.join(self.tmpdir, "deployment.db")
        with patch.object(Config, "DATABASE_PATH", database_path), \
                patch.object(Config, "SEED_DEMO_USERS", False):
            deployment_app = create_app()

        with deployment_app.app_context():
            db = get_db()
            self.assertEqual(db.execute("SELECT COUNT(*) FROM users").fetchone()[0], 0)
            self.assertGreater(db.execute("SELECT COUNT(*) FROM menu_items").fetchone()[0], 0)

            from database.seed_data import ensure_seed_data
            ensure_seed_data()
            self.assertEqual(db.execute("SELECT COUNT(*) FROM users").fetchone()[0], 0)
            self.assertGreater(db.execute("SELECT COUNT(*) FROM menu_items").fetchone()[0], 0)


class TestAccessControl(BaseCase):
    """IA Principle: Confidentiality / least privilege."""

    def test_anonymous_cannot_reach_dashboard(self):
        self.assertEqual(self.client.get("/admin/dashboard").status_code, 302)

    def test_anonymous_cannot_export_transactions(self):
        self.assertEqual(self.client.get("/admin/transactions/export.xlsx").status_code, 302)

    def test_cashier_cannot_reach_dashboard(self):
        self.login_cashier()
        self.assertEqual(self.client.get("/admin/dashboard").status_code, 403)

    def test_cashier_cannot_reach_menu_users_or_audit_log(self):
        self.login_cashier()
        for url in ("/admin/menu", "/admin/users", "/admin/audit-log", "/admin/transactions"):
            self.assertEqual(self.client.get(url).status_code, 403, url)

    def test_cashier_cannot_change_menu_prices(self):
        self.login_cashier()
        res = self.client.post("/admin/menu/update", data={
            "item_id": 1, "price": 1, "cost_price": 1, "csrf_token": self.csrf()})
        self.assertEqual(res.status_code, 403)
        with self.app.app_context():
            price = get_db().execute("SELECT price FROM menu_items WHERE id=1").fetchone()["price"]
        self.assertNotEqual(price, 1)

    def test_cashier_cannot_view_another_cashiers_receipt(self):
        self.login_cashier()
        self.sell()
        self.client.get("/logout")

        self.client.post("/pos-login", data={
            "staff_id": "cashier2", "pin": "2637", "csrf_token": self.csrf(),
        })
        self.assertEqual(self.client.get("/pos/receipt/1").status_code, 403)


class TestInputValidationAndInjection(BaseCase):
    """Secure coding: input validation + parameterized queries."""

    def test_sql_injection_in_login_does_not_authenticate(self):
        res = self.client.post("/admin-login", data={
            "email": "admin' OR '1'='1", "password": "anything", "csrf_token": self.csrf()})
        self.assertIn(b"Invalid username/email or password/PIN", res.data)

    def test_sql_injection_cannot_drop_tables(self):
        self.client.post("/admin-login", data={
            "email": "x'; DROP TABLE users; --", "password": "y", "csrf_token": self.csrf()})
        with self.app.app_context():
            count = get_db().execute("SELECT COUNT(*) AS c FROM users").fetchone()["c"]
        self.assertGreater(count, 0)

    def test_sql_injection_in_menu_name_is_stored_as_text(self):
        self.login_admin()
        payload = "Latte'); DROP TABLE transactions; --"
        self.client.post("/admin/menu/create", data={
            "name": payload, "category": "Coffee Drinks", "price": "100", "cost_price": "40",
            "csrf_token": self.csrf()})
        with self.app.app_context():
            db = get_db()
            db.execute("SELECT COUNT(*) FROM transactions").fetchone()  # table survives
            row = db.execute("SELECT name FROM menu_items WHERE name = ?", (payload,)).fetchone()
        self.assertIsNotNone(row)  # stored as literal text, not executed

    def test_client_cannot_dictate_price(self):
        self.login_cashier()
        self.client.post("/pos/checkout", json={
            "cart": [{"menu_item_id": 1, "quantity": 1, "price": 1}],
            "payment_method": "cash", "service_type": "dine_in",
            "amount_tendered": 1000,
        }, headers={"X-CSRF-Token": self.csrf()})
        with self.app.app_context():
            db = get_db()
            real_price = db.execute("SELECT price FROM menu_items WHERE id=1").fetchone()["price"]
            total = db.execute("SELECT total FROM transactions WHERE id=1").fetchone()["total"]
        self.assertEqual(total, real_price)

    def test_negative_quantity_is_rejected(self):
        self.login_cashier()
        res = self.sell(quantity=-5)
        self.assertEqual(res.status_code, 400)

    def test_invalid_order_type_is_rejected(self):
        self.login_cashier()
        res = self.client.post("/pos/checkout", json={
            "cart": [{"menu_item_id": 1, "quantity": 1}],
            "payment_method": "cash", "service_type": "delivery",
            "amount_tendered": 500,
        }, headers={"X-CSRF-Token": self.csrf()})
        self.assertEqual(res.status_code, 400)
        self.assertIn("Dine In or Take Out", res.get_json()["error"])

    def test_missing_order_type_is_rejected(self):
        self.login_cashier()
        res = self.client.post("/pos/checkout", json={
            "cart": [{"menu_item_id": 1, "quantity": 1}],
            "payment_method": "cash", "amount_tendered": 500,
        }, headers={"X-CSRF-Token": self.csrf()})
        self.assertEqual(res.status_code, 400)

    def test_checkout_records_order_type(self):
        self.login_cashier()
        res = self.client.post("/pos/checkout", json={
            "cart": [{"menu_item_id": 1, "quantity": 1}],
            "payment_method": "cash", "service_type": "takeout",
            "amount_tendered": 500,
        }, headers={"X-CSRF-Token": self.csrf()})
        self.assertEqual(res.status_code, 200)
        with self.app.app_context():
            service_type = get_db().execute(
                "SELECT service_type FROM transactions WHERE id = 1"
            ).fetchone()["service_type"]
        self.assertEqual(service_type, "takeout")

    def test_invalid_payment_method_is_rejected(self):
        self.login_cashier()
        res = self.sell(method="crypto")
        self.assertEqual(res.status_code, 400)

    def test_weak_pin_is_rejected_on_user_creation(self):
        self.login_admin()
        res = self.client.post("/admin/users/create", data={
            "full_name": "Test Person", "username": "testperson", "role": "cashier",
            "pin": "1234", "csrf_token": self.csrf()})
        self.assertIn(b"easily-guessed", res.data)

    def test_cost_above_price_is_rejected(self):
        self.login_admin()
        res = self.client.post("/admin/menu/create", data={
            "name": "Loss Leader", "category": "Coffee Drinks", "price": "50", "cost_price": "80",
            "csrf_token": self.csrf()})
        self.assertIn(b"cannot be higher", res.data)


class TestPayment(BaseCase):
    """POS payment handling — change is computed server-side."""

    def test_cash_change_is_calculated_correctly(self):
        self.login_cashier()
        with self.app.app_context():
            price = get_db().execute("SELECT price FROM menu_items WHERE id=1").fetchone()["price"]
        res = self.sell(tendered=500)
        self.assertEqual(res.status_code, 200)
        self.assertAlmostEqual(res.get_json()["change_due"], 500 - price, places=2)

    def test_insufficient_cash_is_rejected(self):
        self.login_cashier()
        res = self.sell(tendered=1)
        self.assertEqual(res.status_code, 400)
        with self.app.app_context():
            count = get_db().execute("SELECT COUNT(*) AS c FROM transactions").fetchone()["c"]
        self.assertEqual(count, 0)  # no sale written

    def test_client_cannot_fake_the_change_due(self):
        self.login_cashier()
        self.client.post("/pos/checkout", json={
            "cart": [{"menu_item_id": 1, "quantity": 1}],
            "payment_method": "cash", "service_type": "dine_in",
            "amount_tendered": 100, "change_due": 9999,
        }, headers={"X-CSRF-Token": self.csrf()})
        with self.app.app_context():
            db = get_db()
            txn = db.execute("SELECT total, amount_tendered, change_due FROM transactions WHERE id=1").fetchone()
        self.assertAlmostEqual(txn["change_due"], txn["amount_tendered"] - txn["total"], places=2)
        self.assertNotEqual(txn["change_due"], 9999)

    def test_ewallet_payment_records_no_change(self):
        self.login_cashier()
        res = self.sell(method="ewallet", tendered=0)
        self.assertEqual(res.status_code, 200)
        with self.app.app_context():
            txn = get_db().execute("SELECT * FROM transactions WHERE id=1").fetchone()
        self.assertEqual(txn["change_due"], 0)
        self.assertEqual(txn["payment_method"], "ewallet")

    def test_card_payment_is_no_longer_accepted(self):
        self.login_cashier()
        res = self.sell(method="card", tendered=0)
        self.assertEqual(res.status_code, 400)
        with self.app.app_context():
            count = get_db().execute("SELECT COUNT(*) AS c FROM transactions").fetchone()["c"]
        self.assertEqual(count, 0)

    def test_vat_breakdown_is_correct(self):
        self.login_cashier()
        self.sell(menu_item_id=1, quantity=1)
        with self.app.app_context():
            txn = get_db().execute("SELECT total, vat_amount FROM transactions WHERE id=1").fetchone()
        expected_vat = round(txn["total"] - txn["total"] / 1.12, 2)
        self.assertAlmostEqual(txn["vat_amount"], expected_vat, places=2)


class TestCSRF(BaseCase):
    def test_post_without_token_is_blocked(self):
        self.login_cashier()
        res = self.client.post("/pos/checkout", json={
            "cart": [{"menu_item_id": 1, "quantity": 1}], "payment_method": "cash",
            "amount_tendered": 500})
        self.assertEqual(res.status_code, 400)


class TestIntegrityAndNonRepudiation(BaseCase):
    """IA Principles: Integrity + Non-repudiation."""

    def test_void_requires_manager_pin(self):
        self.login_cashier()
        self.sell()
        res = self.client.post("/pos/void/1", json={"manager_pin": "0000", "reason": "test"},
                               headers={"X-CSRF-Token": self.csrf()})
        self.assertEqual(res.status_code, 403)
        with self.app.app_context():
            status = get_db().execute("SELECT status FROM transactions WHERE id=1").fetchone()["status"]
        self.assertEqual(status, "completed")

    def test_void_requires_a_reason(self):
        self.login_cashier()
        self.sell()
        res = self.client.post("/pos/void/1", json={"manager_pin": "7391", "reason": ""},
                               headers={"X-CSRF-Token": self.csrf()})
        self.assertEqual(res.status_code, 400)

    def test_void_with_manager_pin_succeeds_and_records_approver(self):
        self.login_cashier()
        self.sell()
        res = self.client.post("/pos/void/1", json={"manager_pin": "7391", "reason": "wrong order"},
                               headers={"X-CSRF-Token": self.csrf()})
        self.assertEqual(res.status_code, 200)
        with self.app.app_context():
            txn = get_db().execute("SELECT * FROM transactions WHERE id=1").fetchone()
        self.assertEqual(txn["status"], "voided")
        self.assertIsNotNone(txn["voided_by"])
        self.assertEqual(txn["void_reason"], "wrong order")

    def test_voided_sale_is_marked_not_deleted(self):
        self.login_cashier()
        self.sell()
        self.client.post("/pos/void/1", json={"manager_pin": "7391", "reason": "test"},
                         headers={"X-CSRF-Token": self.csrf()})
        with self.app.app_context():
            row = get_db().execute("SELECT * FROM transactions WHERE id=1").fetchone()
        self.assertIsNotNone(row)  # the record still exists — no evidence destroyed

    def test_voided_sale_excluded_from_revenue(self):
        self.login_cashier()
        self.sell()
        self.sell()
        self.client.post("/pos/void/1", json={"manager_pin": "7391", "reason": "test"},
                         headers={"X-CSRF-Token": self.csrf()})
        with self.app.app_context():
            rev = get_db().execute(
                "SELECT COALESCE(SUM(total),0) AS r FROM transactions WHERE status='completed'"
            ).fetchone()["r"]
            price = get_db().execute("SELECT price FROM menu_items WHERE id=1").fetchone()["price"]
        self.assertAlmostEqual(rev, price, places=2)  # only the one surviving sale

    def test_every_sensitive_action_is_logged(self):
        self.login_cashier()
        self.sell()
        self.client.post("/pos/void/1", json={"manager_pin": "7391", "reason": "test"},
                         headers={"X-CSRF-Token": self.csrf()})
        with self.app.app_context():
            actions = [r["action"] for r in get_db().execute("SELECT action FROM audit_log").fetchall()]
        for expected in ("LOGIN_SUCCESS", "SALE_COMPLETED", "VOID_APPROVED"):
            self.assertIn(expected, actions)

    def test_menu_changes_are_logged(self):
        self.login_admin()
        self.client.post("/admin/menu/update", data={
            "item_id": 1, "price": "150", "cost_price": "30", "is_active": "on",
            "csrf_token": self.csrf()})
        with self.app.app_context():
            actions = [r["action"] for r in get_db().execute("SELECT action FROM audit_log").fetchall()]
        self.assertIn("MENU_ITEM_UPDATED", actions)

    def test_receipt_keeps_historical_price_after_menu_change(self):
        self.login_cashier()
        self.sell()
        with self.app.app_context():
            db = get_db()
            original = db.execute("SELECT unit_price FROM transaction_items WHERE transaction_id=1").fetchone()[0]
            db.execute("UPDATE menu_items SET price = 999 WHERE id = 1")
            db.commit()
            still = db.execute("SELECT unit_price FROM transaction_items WHERE transaction_id=1").fetchone()[0]
        self.assertEqual(original, still)  # old receipt unchanged

    def test_profit_uses_cost_recorded_at_sale_time(self):
        self.login_cashier()
        self.sell()
        with self.app.app_context():
            db = get_db()
            before = db.execute(
                "SELECT SUM((unit_price - unit_cost) * quantity) AS p FROM transaction_items"
            ).fetchone()["p"]
            db.execute("UPDATE menu_items SET cost_price = 89 WHERE id = 1")  # cost spikes today
            db.commit()
            after = db.execute(
                "SELECT SUM((unit_price - unit_cost) * quantity) AS p FROM transaction_items"
            ).fetchone()["p"]
        self.assertEqual(before, after)  # yesterday's profit isn't rewritten

    def test_removed_menu_item_cannot_be_sold(self):
        self.login_admin()
        self.client.post("/admin/menu/1/remove", data={"csrf_token": self.csrf()})
        self.client.get("/logout")
        self.login_cashier()
        res = self.sell(menu_item_id=1)
        self.assertEqual(res.status_code, 400)

    def test_removed_menu_item_is_soft_deleted(self):
        self.login_admin()
        self.client.post("/admin/menu/1/remove", data={"csrf_token": self.csrf()})
        with self.app.app_context():
            row = get_db().execute("SELECT is_active FROM menu_items WHERE id=1").fetchone()
        self.assertIsNotNone(row)        # row survives so old receipts resolve
        self.assertEqual(row["is_active"], 0)


class TestAvailabilityOfflineMode(BaseCase):
    """IA Principle: Availability."""

    def test_sales_work_with_internet_completely_down(self):
        import security.connectivity as conn
        original = conn.is_internet_available
        conn.is_internet_available = lambda *a, **k: False  # simulate outage
        try:
            self.login_cashier()
            res = self.sell(quantity=2)
            self.assertEqual(res.status_code, 200)
            self.assertTrue(res.get_json()["success"])
        finally:
            conn.is_internet_available = original

    def test_transactions_queue_as_pending_then_sync(self):
        self.login_cashier()
        self.sell()
        with self.app.app_context():
            status = get_db().execute("SELECT sync_status FROM transactions WHERE id=1").fetchone()[0]
        self.assertEqual(status, "pending")

        self.client.get("/logout")
        self.login_admin()

        import blueprints.sync as syncmod
        original = syncmod.is_internet_available
        syncmod.is_internet_available = lambda *a, **k: True  # simulate reconnect
        try:
            res = self.client.post("/api/sync", headers={"X-CSRF-Token": self.csrf()})
            self.assertEqual(res.status_code, 200)
        finally:
            syncmod.is_internet_available = original

        with self.app.app_context():
            status = get_db().execute("SELECT sync_status FROM transactions WHERE id=1").fetchone()[0]
        self.assertEqual(status, "synced")

    def test_sync_refuses_while_offline_and_keeps_queue(self):
        self.login_cashier()
        self.sell()
        self.client.get("/logout")
        self.login_admin()

        import blueprints.sync as syncmod
        original = syncmod.is_internet_available
        syncmod.is_internet_available = lambda *a, **k: False
        try:
            res = self.client.post("/api/sync", headers={"X-CSRF-Token": self.csrf()})
            self.assertEqual(res.status_code, 503)
        finally:
            syncmod.is_internet_available = original

        with self.app.app_context():
            status = get_db().execute("SELECT sync_status FROM transactions WHERE id=1").fetchone()[0]
        self.assertEqual(status, "pending")  # nothing falsely marked as backed up

    def test_cashier_cannot_trigger_sync(self):
        self.login_cashier()
        res = self.client.post("/api/sync", headers={"X-CSRF-Token": self.csrf()})
        self.assertEqual(res.status_code, 403)


class TestSecurityHeaders(BaseCase):
    def test_headers_present(self):
        res = self.client.get("/")
        self.assertEqual(res.headers.get("X-Content-Type-Options"), "nosniff")
        self.assertEqual(res.headers.get("X-Frame-Options"), "DENY")
        self.assertIn("default-src 'self'", res.headers.get("Content-Security-Policy", ""))


class TestBootstrapAdminSetup(BaseCase):
    """First-run owner setup: the only login-page path that creates an
    account, and it can only ever create an admin. seed_admin=False gives
    each test in this class a database with zero admin accounts, matching
    a genuine fresh install."""
    seed_admin = False

    def valid_setup_data(self, **overrides):
        data = {
            "full_name": "Cafe Owner", "email": "owner@example.com",
            "password": "Owner#2026", "pin": "5170",
            "csrf_token": self.csrf(),
        }
        data.update(overrides)
        return data

    def test_admin_login_redirects_to_setup_when_no_admin_exists(self):
        res = self.client.get("/admin-login")
        self.assertEqual(res.status_code, 302)
        self.assertIn("/admin-setup", res.headers["Location"])

    def test_setup_form_is_shown_when_no_admin_exists(self):
        res = self.client.get("/admin-setup")
        self.assertEqual(res.status_code, 200)
        self.assertIn(b"Create Owner Account", res.data)

    def test_first_admin_setup_requires_configured_deployment_token(self):
        self.app.config["ADMIN_SETUP_TOKEN"] = "private-one-time-setup-token"
        form = self.client.get("/admin-setup")
        self.assertIn(b"One-time deployment setup token", form.data)

        rejected = self.client.post("/admin-setup", data=self.valid_setup_data())
        self.assertIn(b"Setup authorization failed", rejected.data)
        with self.app.app_context():
            count = get_db().execute(
                "SELECT COUNT(*) FROM users WHERE role='admin'"
            ).fetchone()[0]
        self.assertEqual(count, 0)

        accepted = self.client.post("/admin-setup", data=self.valid_setup_data(
            setup_token="private-one-time-setup-token",
        ))
        self.assertEqual(accepted.status_code, 302)
        with self.app.app_context():
            count = get_db().execute(
                "SELECT COUNT(*) FROM users WHERE role='admin'"
            ).fetchone()[0]
        self.assertEqual(count, 1)

    def test_setup_creates_an_admin_and_logs_them_in(self):
        res = self.client.post("/admin-setup", data=self.valid_setup_data(), follow_redirects=True)
        self.assertEqual(res.status_code, 200)
        self.assertIn(b"Owner Dashboard", res.data)  # landed on the dashboard, already logged in
        with self.app.app_context():
            row = get_db().execute(
                "SELECT role, full_name, username FROM users WHERE email='owner@example.com'"
            ).fetchone()
        self.assertEqual(row["role"], "admin")
        self.assertEqual(row["full_name"], "Cafe Owner")
        self.assertEqual(row["username"], "owner@example.com")  # derived from email

    def test_setup_never_creates_a_cashier_role(self):
        # The form has no role field at all, but confirm the route can't be
        # tricked into writing a different role even if one is posted.
        data = self.valid_setup_data()
        data["role"] = "cashier"
        self.client.post("/admin-setup", data=data)
        with self.app.app_context():
            role = get_db().execute("SELECT role FROM users WHERE email='owner@example.com'").fetchone()["role"]
        self.assertEqual(role, "admin")

    def test_setup_closes_after_first_admin_is_created(self):
        self.client.post("/admin-setup", data=self.valid_setup_data())
        self.client.get("/logout")

        res = self.client.get("/admin-setup")
        self.assertEqual(res.status_code, 302)
        self.assertIn("/admin-login", res.headers["Location"])

    def test_cannot_create_a_second_admin_through_setup(self):
        self.client.post("/admin-setup", data=self.valid_setup_data())
        self.client.get("/logout")

        res = self.client.post("/admin-setup", data=self.valid_setup_data(
            email="second-owner@example.com", csrf_token=self.csrf(),
        ))
        self.assertEqual(res.status_code, 302)
        with self.app.app_context():
            count = get_db().execute(
                "SELECT COUNT(*) AS c FROM users WHERE role='admin'"
            ).fetchone()["c"]
        self.assertEqual(count, 1)

    def test_weak_pin_rejected_on_setup(self):
        res = self.client.post("/admin-setup", data=self.valid_setup_data(pin="0000"))
        self.assertIn(b"easily-guessed", res.data)
        with self.app.app_context():
            count = get_db().execute("SELECT COUNT(*) AS c FROM users WHERE role='admin'").fetchone()["c"]
        self.assertEqual(count, 0)

    def test_weak_password_rejected_on_setup(self):
        res = self.client.post("/admin-setup", data=self.valid_setup_data(password="short"))
        self.assertIn(b"Password must be at least", res.data)

    def test_invalid_email_rejected_on_setup(self):
        res = self.client.post("/admin-setup", data=self.valid_setup_data(email="not-an-email"))
        self.assertIn(b"valid email", res.data)
        with self.app.app_context():
            count = get_db().execute("SELECT COUNT(*) AS c FROM users WHERE role='admin'").fetchone()["c"]
        self.assertEqual(count, 0)

    def test_setup_creation_is_logged(self):
        self.client.post("/admin-setup", data=self.valid_setup_data())
        with self.app.app_context():
            actions = [r["action"] for r in get_db().execute("SELECT action FROM audit_log").fetchall()]
        self.assertIn("USER_CREATED", actions)

    def test_created_owner_can_log_in_normally_afterward(self):
        self.client.post("/admin-setup", data=self.valid_setup_data())
        self.client.get("/logout")
        res = self.client.post("/admin-login", data={
            "email": "owner@example.com", "password": "Owner#2026", "csrf_token": self.csrf(),
        }, follow_redirects=True)
        self.assertIn(b"Owner Dashboard", res.data)


class TestStaffPanelAdminCreation(BaseCase):
    """Creating a SECOND admin from inside the back office (as opposed to
    the first-run bootstrap) needs an email too, and for the same reason:
    without one, that account could never use forgot-password."""

    def create_admin(self, **overrides):
        data = {
            "full_name": "Second Manager", "role": "admin",
            "email": "manager2@example.com", "password": "Manager#26", "pin": "8420",
            "csrf_token": self.csrf(),
        }
        data.update(overrides)
        return self.client.post("/admin/users/create", data=data)

    def test_creating_admin_requires_valid_email(self):
        self.login_admin()
        res = self.create_admin(email="not-an-email")
        self.assertIn(b"valid email", res.data)

    def test_creating_admin_requires_a_password(self):
        self.login_admin()
        res = self.create_admin(password="")
        self.assertIn(b"Admin accounts need a password", res.data)

    def test_new_admin_can_log_in_and_reach_dashboard(self):
        self.login_admin()
        self.create_admin()
        self.client.get("/logout")

        res = self.client.post("/admin-login", data={
            "email": "manager2@example.com", "password": "Manager#26", "csrf_token": self.csrf(),
        }, follow_redirects=True)
        self.assertIn(b"Owner Dashboard", res.data)

    def test_duplicate_admin_email_is_rejected(self):
        self.login_admin()
        res = self.create_admin(email="admin@example.com")  # already the fixture admin
        self.assertIn(b"already exists", res.data)

    def test_cashier_creation_still_uses_username_not_email(self):
        self.login_admin()
        res = self.client.post("/admin/users/create", data={
            "full_name": "New Cashier", "role": "cashier",
            "username": "newcashier", "pin": "3344", "csrf_token": self.csrf(),
        })
        self.assertEqual(res.status_code, 302)
        with self.app.app_context():
            row = get_db().execute("SELECT email, username FROM users WHERE username='newcashier'").fetchone()
        self.assertIsNone(row["email"])
        self.assertEqual(row["username"], "newcashier")


class TestForgotPasswordOTP(BaseCase):
    """Authentication (recovery) + Availability of the account itself.

    send_otp_email is monkey-patched on blueprints.auth (where the route
    actually calls it, not security.mailer where it's defined — `from x
    import y` binds a new name in the importing module) so no real
    SMTP/network call happens. The patch captures what WOULD have been
    sent so a test can read the OTP back out, the same way a real person
    would read it from their inbox, and use it in a follow-up request.
    """

    def setUp(self):
        super().setUp()
        self._env_backup = {
            "GMAIL_ADDRESS": os.environ.get("GMAIL_ADDRESS"),
            "GMAIL_APP_PASSWORD": os.environ.get("GMAIL_APP_PASSWORD"),
        }
        os.environ["GMAIL_ADDRESS"] = "test-sender@example.com"
        os.environ["GMAIL_APP_PASSWORD"] = "fake-app-password-for-tests"

        import blueprints.auth as auth_module
        self._auth_module = auth_module
        self._original_send = auth_module.send_otp_email
        self.sent = []
        auth_module.send_otp_email = lambda to, otp, name: self.sent.append(
            {"to": to, "otp": otp, "name": name})

    def tearDown(self):
        self._auth_module.send_otp_email = self._original_send
        for key, value in self._env_backup.items():
            if value is None:
                os.environ.pop(key, None)
            else:
                os.environ[key] = value
        super().tearDown()

    def request_reset(self, email="admin@example.com"):
        return self.client.post("/forgot-password", data={"email": email, "csrf_token": self.csrf()})

    def test_request_for_existing_admin_sends_a_six_digit_otp(self):
        res = self.request_reset()
        self.assertIn(b"reset code has been sent", res.data)
        self.assertEqual(len(self.sent), 1)
        self.assertEqual(self.sent[0]["to"], "admin@example.com")
        self.assertEqual(len(self.sent[0]["otp"]), 6)
        self.assertTrue(self.sent[0]["otp"].isdigit())

    def test_request_for_nonexistent_email_shows_same_message_but_sends_nothing(self):
        res = self.request_reset(email="nobody@example.com")
        self.assertIn(b"reset code has been sent", res.data)
        self.assertEqual(len(self.sent), 0)

    def test_unconfigured_email_gives_identical_response_regardless_of_account(self):
        os.environ.pop("GMAIL_ADDRESS", None)
        os.environ.pop("GMAIL_APP_PASSWORD", None)
        real = self.request_reset(email="admin@example.com").data
        fake = self.request_reset(email="nobody@example.com").data
        self.assertIn(b"configured on this server yet", real)
        self.assertEqual(real, fake)
        self.assertEqual(len(self.sent), 0)

    def test_full_round_trip_changes_password_and_logs_in(self):
        self.request_reset()
        otp = self.sent[0]["otp"]
        res = self.client.post("/reset-password", data={
            "email": "admin@example.com", "otp": otp, "new_password": "NewPass#2026",
            "csrf_token": self.csrf(),
        }, follow_redirects=True)
        self.assertIn(b"Owner Dashboard", res.data)

        self.client.get("/logout")
        res = self.client.post("/admin-login", data={
            "email": "admin@example.com", "password": "NewPass#2026", "csrf_token": self.csrf(),
        }, follow_redirects=True)
        self.assertIn(b"Owner Dashboard", res.data)

    def test_wrong_otp_is_rejected_and_old_password_still_works(self):
        self.request_reset()
        res = self.client.post("/reset-password", data={
            "email": "admin@example.com", "otp": "000000", "new_password": "NewPass#2026",
            "csrf_token": self.csrf(),
        })
        self.assertIn(b"invalid or has expired", res.data)

        res = self.client.post("/admin-login", data={
            "email": "admin@example.com", "password": "Admin#2026", "csrf_token": self.csrf(),
        }, follow_redirects=True)
        self.assertIn(b"Owner Dashboard", res.data)

    def test_otp_is_single_use(self):
        self.request_reset()
        otp = self.sent[0]["otp"]
        self.client.post("/reset-password", data={
            "email": "admin@example.com", "otp": otp, "new_password": "FirstPass#26",
            "csrf_token": self.csrf(),
        })
        self.client.get("/logout")
        res = self.client.post("/reset-password", data={
            "email": "admin@example.com", "otp": otp, "new_password": "SecondPass#26",
            "csrf_token": self.csrf(),
        })
        self.assertIn(b"invalid or has expired", res.data)

    def test_expired_otp_is_rejected(self):
        self.request_reset()
        with self.app.app_context():
            db = get_db()
            db.execute("UPDATE password_resets SET expires_at = '2000-01-01 00:00:00'")
            db.commit()
        otp = self.sent[0]["otp"]
        res = self.client.post("/reset-password", data={
            "email": "admin@example.com", "otp": otp, "new_password": "NewPass#2026",
            "csrf_token": self.csrf(),
        })
        self.assertIn(b"invalid or has expired", res.data)

    def test_too_many_wrong_attempts_locks_out_even_the_correct_code(self):
        self.request_reset()
        otp = self.sent[0]["otp"]
        for _ in range(5):
            self.client.post("/reset-password", data={
                "email": "admin@example.com", "otp": "111111", "new_password": "NewPass#2026",
                "csrf_token": self.csrf(),
            })
        res = self.client.post("/reset-password", data={
            "email": "admin@example.com", "otp": otp, "new_password": "NewPass#2026",
            "csrf_token": self.csrf(),
        })
        self.assertIn(b"invalid or has expired", res.data)

    def test_new_request_invalidates_the_previous_code(self):
        self.request_reset()
        first_otp = self.sent[0]["otp"]
        with self.app.app_context():
            db = get_db()
            db.execute("UPDATE password_resets SET created_at = '2000-01-01 00:00:00'")  # skip cooldown
            db.commit()
        self.request_reset()
        second_otp = self.sent[1]["otp"]
        self.assertNotEqual(first_otp, second_otp)

        res = self.client.post("/reset-password", data={
            "email": "admin@example.com", "otp": first_otp, "new_password": "NewPass#2026",
            "csrf_token": self.csrf(),
        })
        self.assertIn(b"invalid or has expired", res.data)

    def test_resend_cooldown_prevents_immediate_second_send(self):
        self.request_reset()
        self.request_reset()
        self.assertEqual(len(self.sent), 1)

    def test_weak_new_password_rejected(self):
        self.request_reset()
        otp = self.sent[0]["otp"]
        res = self.client.post("/reset-password", data={
            "email": "admin@example.com", "otp": otp, "new_password": "weak",
            "csrf_token": self.csrf(),
        })
        self.assertIn(b"New password must be at least", res.data)

    def test_reset_flow_actions_are_logged(self):
        self.request_reset()
        otp = self.sent[0]["otp"]
        self.client.post("/reset-password", data={
            "email": "admin@example.com", "otp": "000000", "new_password": "NewPass#2026",
            "csrf_token": self.csrf(),
        })
        self.client.post("/reset-password", data={
            "email": "admin@example.com", "otp": otp, "new_password": "NewPass#2026",
            "csrf_token": self.csrf(),
        })
        with self.app.app_context():
            actions = [r["action"] for r in get_db().execute("SELECT action FROM audit_log").fetchall()]
        self.assertIn("PASSWORD_RESET_REQUESTED", actions)
        self.assertIn("PASSWORD_RESET_FAILED", actions)
        self.assertIn("PASSWORD_RESET_COMPLETED", actions)


class TestSchemaMigration(unittest.TestCase):
    """A database created by a version of this app before the email column
    existed must be upgraded in place on next startup, without losing any
    existing rows — see database/db.py's run_migrations()."""

    def test_old_database_gains_email_column_without_losing_data(self):
        import sqlite3

        tmpdir = tempfile.mkdtemp()
        db_path = os.path.join(tmpdir, "old.db")

        # Simulate a database file left behind by the pre-email version of
        # this app: same users table, but no email column at all.
        old_conn = sqlite3.connect(db_path)
        old_conn.execute("""
            CREATE TABLE users (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                username TEXT UNIQUE NOT NULL,
                password_hash TEXT,
                pin_hash TEXT NOT NULL,
                full_name TEXT NOT NULL,
                role TEXT NOT NULL,
                is_active INTEGER NOT NULL DEFAULT 1,
                failed_attempts INTEGER NOT NULL DEFAULT 0,
                locked_until TEXT,
                created_at TEXT NOT NULL
            )
        """)
        old_conn.execute(
            "INSERT INTO users (username, pin_hash, full_name, role, is_active, created_at) "
            "VALUES ('legacyadmin', 'somehash', 'Legacy Admin', 'admin', 1, '2026-01-01 00:00:00')"
        )
        old_conn.commit()
        old_conn.close()

        os.environ["SECRET_KEY"] = "test-secret-key-not-for-production"
        app = create_app()
        app.config["TESTING"] = True
        app.config["DATABASE_PATH"] = db_path

        with app.app_context():
            from database.seed_data import ensure_seed_data, init_schema
            init_schema(app)       # must not crash on the pre-existing table, must add the column
            ensure_seed_data()     # must not re-seed — real data (however old) already exists
            db = get_db()
            row = db.execute("SELECT * FROM users WHERE username='legacyadmin'").fetchone()
            columns = [r["name"] for r in db.execute("PRAGMA table_info(users)").fetchall()]

        self.assertIn("email", columns)
        self.assertIsNotNone(row)              # the pre-existing row is untouched
        self.assertEqual(row["full_name"], "Legacy Admin")
        self.assertIsNone(row["email"])        # new column, simply empty — no data invented

    def test_migration_is_safe_to_run_twice(self):
        os.environ["SECRET_KEY"] = "test-secret-key-not-for-production"
        tmpdir = tempfile.mkdtemp()
        app = create_app()
        app.config["TESTING"] = True
        app.config["DATABASE_PATH"] = os.path.join(tmpdir, "fresh.db")

        with app.app_context():
            from database.seed_data import ensure_seed_data, init_schema
            init_schema(app)
            init_schema(app)  # run again — must not raise "duplicate column"
            ensure_seed_data()
            columns = [r["name"] for r in get_db().execute("PRAGMA table_info(users)").fetchall()]
        self.assertIn("email", columns)

    def test_historical_card_transactions_still_display_after_card_was_removed(self):
        # Builds a database matching what existed before 'card' was dropped
        # from the payment_method CHECK constraint — its own users and
        # transactions tables already in place, one of them a card sale —
        # then starts the real app against that file, the way upgrading an
        # actual deployment would. CREATE TABLE IF NOT EXISTS leaves both
        # tables untouched (see schema.sql's comment on this), so the old,
        # more permissive constraint — and the card row it allowed — should
        # both survive completely unaffected.
        import sqlite3

        from werkzeug.security import generate_password_hash

        tmpdir = tempfile.mkdtemp()
        db_path = os.path.join(tmpdir, "pre_card_removal.db")

        old_conn = sqlite3.connect(db_path)
        old_conn.execute("""
            CREATE TABLE users (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                username TEXT UNIQUE NOT NULL,
                email TEXT,
                password_hash TEXT,
                pin_hash TEXT NOT NULL,
                full_name TEXT NOT NULL,
                role TEXT NOT NULL,
                is_active INTEGER NOT NULL DEFAULT 1,
                failed_attempts INTEGER NOT NULL DEFAULT 0,
                locked_until TEXT,
                created_at TEXT NOT NULL
            )
        """)
        old_conn.execute("""
            CREATE TABLE transactions (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                cashier_id INTEGER NOT NULL REFERENCES users(id),
                created_at TEXT NOT NULL,
                subtotal REAL NOT NULL, vat_amount REAL NOT NULL, total REAL NOT NULL,
                payment_method TEXT NOT NULL DEFAULT 'cash'
                    CHECK (payment_method IN ('cash', 'card', 'ewallet')),
                amount_tendered REAL NOT NULL DEFAULT 0, change_due REAL NOT NULL DEFAULT 0,
                status TEXT NOT NULL DEFAULT 'completed' CHECK (status IN ('completed', 'voided')),
                voided_by INTEGER REFERENCES users(id), voided_at TEXT, void_reason TEXT,
                sync_status TEXT NOT NULL DEFAULT 'pending' CHECK (sync_status IN ('pending', 'synced')),
                synced_at TEXT
            )
        """)
        old_conn.execute(
            "INSERT INTO users (username, email, password_hash, pin_hash, full_name, role, "
            "is_active, created_at) VALUES ('owner@example.com', 'owner@example.com', ?, ?, "
            "'Legacy Owner', 'admin', 1, '2026-01-01 00:00:00')",
            (generate_password_hash("Legacy#2026"), generate_password_hash("1357")),
        )
        old_conn.execute(
            "INSERT INTO transactions (cashier_id, created_at, subtotal, vat_amount, total, "
            "payment_method, amount_tendered, change_due, status, sync_status) "
            "VALUES (1, '2026-06-01 10:00:00', 100, 10.71, 100, 'card', 100, 0, 'completed', 'synced')"
        )
        old_conn.commit()
        old_conn.close()

        os.environ["SECRET_KEY"] = "test-secret-key-not-for-production"
        app = create_app()
        app.config["TESTING"] = True
        app.config["DATABASE_PATH"] = db_path
        app.test_client_class = CaptchaTestClient
        client = app.test_client()

        with app.app_context():
            from database.seed_data import ensure_seed_data, init_schema
            init_schema(app)       # users/transactions already exist — must not touch their constraints
            ensure_seed_data()     # users table is non-empty — must not re-seed over the legacy admin
            db = get_db()
            constraint_sql = db.execute(
                "SELECT sql FROM sqlite_master WHERE name='transactions'"
            ).fetchone()["sql"]
            migrated_txn = db.execute(
                "SELECT service_type, is_demo FROM transactions WHERE id=1"
            ).fetchone()
        self.assertIn("'card'", constraint_sql)  # old constraint, untouched
        self.assertEqual(migrated_txn["service_type"], "takeout")
        self.assertEqual(migrated_txn["is_demo"], 0)

        token = client.get("/api/csrf-token").get_json()["token"]
        client.post("/admin-login", data={
            "email": "owner@example.com", "password": "Legacy#2026", "csrf_token": token})

        dashboard = client.get("/admin/dashboard").data
        self.assertIn(b"Card", dashboard)  # Recent Transactions renders it via |title, doesn't crash

        all_txns = client.get("/admin/transactions").data
        self.assertIn(b"Card", all_txns)


if __name__ == "__main__":
    unittest.main(verbosity=2)
